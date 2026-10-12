"""The protocol as a .docx somebody can print, sign and take to the bench.

A Word file is a zip of XML, and the failure mode of writing one by hand is
not a wrong heading — it is a file Word refuses to open at all, with no useful
message, usually because a quote from a paper contained an ampersand. So these
tests open the bytes rather than trusting them: every part is parsed as XML,
and where a real OOXML reader is installed the document is read back through
it.

The rest is about what the document says. A protocol this system wrote is a
proposal, and a document that reads like a work order is the most dangerous
thing it could produce — so "not approved" is checked, the open values are
checked to come before the steps, and provenance is checked to survive into
the page.
"""
import io
import unittest
import xml.dom.minidom
import zipfile

from biosense.production import lab_protocol as LP
from biosense.production.docx import Document

PROTOCOL = {
    'title': 'iPSC to macrophage, yield',
    'objective': 'increase viable macrophage yield from iPSC',
    'created_at': '2026-10-12T00:00:00Z',
    'project': {'project_id': 'ipsc_macrophage', 'name': 'iPSC → macrophage'},
    'evidence_status': 'A real run: the numbers come from the literature, not from a model.',
    'timeline': {'total_days': 25, 'candidates': [],
                 'stages': [{'stage_id': 'expansion', 'label': 'Expansion', 'day_start': 0,
                             'days': 4},
                            {'stage_id': 'myeloid', 'label': 'Myeloid harvest', 'day_start': 4,
                             'days': 21}]},
    'stages': [{'stage_id': 'myeloid', 'parameters': [
        {'parameter_id': 'mcsf_ng_ml', 'label': 'M-CSF & GM-CSF <mix>', 'unit': 'ng/mL',
         'recommended_value': 100, 'control_value': 50, 'changed': True,
         'provenance': 'adapted', 'hypothesis_ref': 'H01'},
        {'parameter_id': 'agitation_rpm', 'label': 'Agitation', 'unit': 'rpm',
         'recommended_value': 50, 'control_value': 50, 'changed': False,
         'provenance': 'design_choice'}]}],
    'hypothesis_ledger': [
        {'hypothesis_id': 'H01', 'adopted': True, 'status': 'proposed', 'confidence': 'moderate',
         'statement': 'Raising M-CSF through the harvest phase raises monocyte output.',
         'lever': {'parameter_id': 'mcsf_ng_ml', 'label': 'M-CSF', 'direction': 'increase',
                   'candidate_value': 100, 'unit': 'ng/mL', 'stage': 'myeloid'}}],
    'readouts': [{'readout_id': 'yield', 'label': 'Monocytes per input cell',
                  'unit': 'cells/input_cell', 'kind': 'yield'}],
    'gaps': [{'parameter_id': 'gmcsf_ng_ml', 'why': 'no source doses it for a stirred tank'}],
    'blocks_wet_lab': True,
    'limitations': ['The simulator does not model this parameter.'],
    'round_plan': {'round': 1, 'replicates': 3,
                   'arms': [{'arm_id': 'control', 'control': True, 'setpoints': {}},
                            {'arm_id': 'mcsf-high', 'setpoints': {'mcsf_ng_ml': 100}}],
                   'unknowns': [{'unknown_id': 'U1', 'question': 'Does M-CSF above 50 help?',
                                 'why_unresolved': 'no source reports it for this vessel'}]},
}


def text_of(data):
    """Every w:t in the document part, in order."""
    z = zipfile.ZipFile(io.BytesIO(data))
    dom = xml.dom.minidom.parseString(z.read('word/document.xml'))
    return [n.firstChild.nodeValue if n.firstChild else ''
            for n in dom.getElementsByTagName('w:t')]


class WriterTests(unittest.TestCase):
    def test_every_part_of_the_file_is_well_formed_xml(self):
        """The failure this guards is not a wrong heading; it is Word refusing
        to open the file at all."""
        d = Document(title='t').heading('H').para('body').bullet('b').table(['A'], [['1']])
        z = zipfile.ZipFile(io.BytesIO(d.tobytes()))
        self.assertIn('word/document.xml', z.namelist())
        self.assertIn('[Content_Types].xml', z.namelist())
        for name in z.namelist():
            xml.dom.minidom.parseString(z.read(name))

    def test_characters_a_paper_quote_contains_do_not_break_the_file(self):
        d = Document().para('IL-3 & M-CSF <100 ng/mL> "quoted"')
        data = d.tobytes()
        xml.dom.minidom.parseString(zipfile.ZipFile(io.BytesIO(data)).read('word/document.xml'))
        self.assertIn('IL-3 & M-CSF <100 ng/mL> "quoted"', text_of(data))

    def test_a_control_character_is_dropped_rather_than_written(self):
        """XML 1.0 has no way to carry one, and a scraped quote sometimes has
        one in it."""
        data = Document().para('clean\x07text').tobytes()
        xml.dom.minidom.parseString(zipfile.ZipFile(io.BytesIO(data)).read('word/document.xml'))
        self.assertIn('cleantext', text_of(data))

    def test_the_same_document_twice_is_the_same_bytes(self):
        """So a diff of two exports shows what changed in the protocol, not
        what time it was written."""
        self.assertEqual(Document(title='x').para('a').tobytes(),
                         Document(title='x').para('a').tobytes())

    def test_a_real_reader_opens_it(self):
        d = Document(title='Protocol')
        d.title_block('Lab protocol', 'draft').heading('Expansion', 1)
        d.para(runs=[('Add ', {}), ('M-CSF', {'bold': True})])
        d.bullet('one').step('two').table(['F', 'D'], [['M-CSF', '100']])
        try:
            import docx
        except ImportError:
            self.skipTest('python-docx is not installed; the XML checks above still ran')
        got = docx.Document(io.BytesIO(d.tobytes()))
        self.assertEqual(['Title', 'Caveat', 'Heading 1', 'Normal', 'List Bullet',
                          'List Number', 'Normal'],
                         [p.style.name for p in got.paragraphs])
        self.assertEqual('Add M-CSF', got.paragraphs[3].text)
        self.assertEqual([['F', 'D'], ['M-CSF', '100']],
                         [[c.text for c in r.cells] for r in got.tables[0].rows])


class ProtocolDocumentTests(unittest.TestCase):
    def setUp(self):
        self.text = text_of(LP.tobytes(PROTOCOL, run_id='ai-20261012-abc123'))
        self.joined = '\n'.join(self.text)

    def test_it_says_it_is_not_approved_at_the_top_and_at_the_bottom(self):
        """A document that reads like a work order is the most dangerous thing
        this system could print."""
        hits = [i for i, t in enumerate(self.text) if 'NOT APPROVED' in t]
        self.assertEqual(2, len(hits))
        self.assertLess(hits[0], len(self.text) / 2, 'the first one is near the top')
        self.assertIn('approve-protocol', self.joined)
        self.assertIn('Nothing in BioSense can do it for them', self.joined)

    def test_what_is_undecided_comes_before_what_to_do(self):
        """Finding a gap at the bench on the day is the failure this ordering
        exists to prevent."""
        gap = self.text.index('Before you run this: what is not decided')
        steps = self.text.index('The protocol, stage by stage')
        self.assertLess(gap, steps)
        self.assertIn('gmcsf_ng_ml', self.joined)
        self.assertIn('no source doses it for a stirred tank', self.joined)

    def test_every_stage_of_the_process_has_its_own_section(self):
        self.assertIn('Expansion — day 0 to 4 (4 d)', self.joined)
        self.assertIn('Myeloid harvest — day 4 to 25 (21 d)', self.joined)

    def test_a_stage_with_nothing_in_it_says_why_rather_than_being_blank(self):
        self.assertIn('No factor is set for this stage by this run', self.joined)
        self.assertIn('not a statement that the stage needs none', self.joined.replace('\n', ' '))

    def test_a_dose_carries_what_it_was_and_what_stands_behind_it(self):
        self.assertIn('100 ng/mL (was 50)', self.text,
                      'the previous value travels with the new one')
        self.assertIn('A', self.text)
        self.assertIn('A — adapted: a source states it for a different system and it was '
                      'moved here.', self.text)

    def test_the_vessel_settings_are_their_own_table_for_whoever_enters_them(self):
        self.assertIn('Bioreactor settings', self.text)
        self.assertIn('Agitation', self.text)
        self.assertIn('50 rpm', self.text)

    def test_the_round_is_laid_out_one_vessel_per_row(self):
        self.assertIn('V1', self.text)
        self.assertIn('V6', self.text, '2 arms × 3 replicates')
        self.assertNotIn('V7', self.text)
        self.assertIn('control (control)', self.joined)

    def test_the_approval_block_is_a_table_to_sign(self):
        for row in ('Reviewed by', 'Date', 'Signature', 'Changes made at the bench'):
            self.assertIn(row, self.text)

    def test_it_never_claims_the_protocol_works(self):
        for word in ('optimal', 'proven', 'guarantee', 'will increase', 'best'):
            self.assertNotIn(word, self.joined.lower(), word)

    def test_a_protocol_with_no_hypothesis_still_produces_a_usable_document(self):
        thin = dict(PROTOCOL, hypothesis_ledger=[], round_plan=None, gaps=[], readouts=[])
        t = '\n'.join(text_of(LP.tobytes(thin)))
        self.assertIn('adopted no hypothesis', t)
        self.assertIn('The protocol, stage by stage', t)
        self.assertIn('NOT APPROVED', t)

    def test_an_empty_protocol_does_not_crash_the_export(self):
        t = '\n'.join(text_of(LP.tobytes({})))
        self.assertIn('NOT APPROVED', t)


class BaselineInTheDocumentTests(unittest.TestCase):
    """A published factor reaches the bench copy marked as published."""

    def test_a_published_factor_appears_with_its_reference_and_no_invented_dose(self):
        from biosense import projects as PJ
        from biosense.evidence import baseline as BL
        from biosense.production import first_pass as FP
        base = BL.build(BL.TEMPLATE, project=PJ.load('ipsc_macrophage'))
        p = dict(PROTOCOL)
        p['first_pass'] = FP.plan(p, baseline=base)
        t = '\n'.join(text_of(LP.tobytes(p)))
        self.assertIn('Y-27632 (ROCK inhibitor)', t)
        self.assertIn('Watanabe2007', t)
        self.assertIn('first 24 h after seeding only', t)
        self.assertIn('GM-CSF', t)
        self.assertIn('no dose published', t)
        self.assertIn('PUB', t)


class DownloadTests(unittest.TestCase):
    """Through the server: the person clicks a link and gets a file."""

    @classmethod
    def setUpClass(cls):
        import json
        import os
        import shutil
        import tempfile
        import threading
        from http.server import ThreadingHTTPServer
        from pathlib import Path
        from biosense import contracts as K
        from biosense import workspace as WS
        from biosense.production import app as APP
        from biosense.production import authz as AZ
        from biosense.production import budget as BU
        from biosense.production import run_store as RS
        from biosense.production import runtime as RT
        cls.tmp = Path(tempfile.mkdtemp())
        cls._rm = lambda: shutil.rmtree(cls.tmp, ignore_errors=True)
        (cls.tmp / 'runs').mkdir()
        cls._private = os.environ.get('BIOSENSE_PRIVATE_DATA')
        os.environ['BIOSENSE_PRIVATE_DATA'] = str(cls.tmp / 'private')
        cfg = RT.from_env(runs_dir=cls.tmp / 'runs', env={})
        APP.Handler.runs_dir = cls.tmp / 'runs'
        APP.Handler.static_dir = K.ROOT / 'webapp'
        APP.Handler.registry = APP.Registry(cls.tmp / 'runs')
        APP.Handler.runtime_cfg = cfg
        APP.Handler.limits = BU.Limits()
        APP.Handler.policy = AZ.Policy()
        APP.Handler.discovery = APP.DiscoveryRegistry(cls.tmp / 'runs', cfg)
        APP.Handler.sessions = WS.SessionStore()
        APP.Handler.default_identity = WS.local_identity()
        APP.Handler._ready = {'at': 0.0, 'body': None, 'code': 503}
        cls.rid = 'f' * 16
        d = cls.tmp / 'runs' / 'ai-20261012-ffffff'
        d.mkdir()
        (d / 'discovery_request.json').write_text(json.dumps(
            {'project_id': 'ipsc_macrophage', 'runtime_mode': 'local_real_ai',
             'objective': 'o', 'request_id': 'r1'}))
        RS.write_state(d, {'run_id': cls.rid, 'status': 'done', 'is_real': True,
                           'runtime_mode': 'local_real_ai', 'owner': None,
                           'project_id': 'ipsc_macrophage', 'objective': 'o',
                           'run_dir': d.name, 'started_at': 1.0, 'finished_at': 2.0,
                           'result': {'protocol': PROTOCOL}})
        cls.srv = ThreadingHTTPServer(('127.0.0.1', 0), APP.Handler)
        cls.srv.daemon_threads = True
        cls.port = cls.srv.server_address[1]
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        import os
        cls.srv.shutdown(); cls.srv.server_close()
        if cls._private is None:
            os.environ.pop('BIOSENSE_PRIVATE_DATA', None)
        else:
            os.environ['BIOSENSE_PRIVATE_DATA'] = cls._private
        cls._rm()

    def test_the_link_returns_a_word_file_named_after_the_run(self):
        import urllib.request
        url = f'http://127.0.0.1:{self.port}/api/discovery/{self.rid}/protocol?format=docx'
        with urllib.request.urlopen(url, timeout=30) as r:
            body = r.read()
            self.assertEqual('application/vnd.openxmlformats-officedocument.'
                             'wordprocessingml.document', r.headers['Content-Type'])
            self.assertIn('attachment; filename=', r.headers['Content-Disposition'])
            self.assertIn('.docx', r.headers['Content-Disposition'])
        self.assertTrue(body.startswith(b'PK'), 'a .docx is a zip')
        self.assertIn('NOT APPROVED', '\n'.join(text_of(body)))

    def test_the_page_offers_the_download(self):
        from biosense import contracts as K
        js = (K.ROOT / 'webapp' / 'discovery.js').read_text()
        self.assertIn('protocol?format=docx', js)
        self.assertIn('Protocol for the bench (Word)', js)


if __name__ == '__main__':
    unittest.main()
